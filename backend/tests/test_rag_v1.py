"""Phase 4 RAG v1: document parsing, chunking, upload, sqlite-vec, citations."""

import io
import json
import uuid
from uuid import UUID

import pytest
from fastapi.testclient import TestClient

from app.db import models
from app.db.database import SessionLocal
from app.main import app
from app.services.documents import chunk_text, parse_document
from app.services.embeddings import _hashing_vector
from app.services.node_handlers import _make_kb_retrieve_fn

from conftest import valid_graph

client = TestClient(app)


def _vec_supported() -> bool:
    """sqlite-vec needs a Python whose sqlite3 exposes loadable extensions —
    absent on some Linux/CI builds. Skip rather than error there."""
    import sqlite3
    if not hasattr(sqlite3.Connection, "enable_load_extension"):
        return False
    try:
        import sqlite_vec
        conn = sqlite3.connect(":memory:")
        conn.enable_load_extension(True)
        sqlite_vec.load(conn)
        conn.close()
        return True
    except Exception:
        return False


vec_required = pytest.mark.skipif(
    not _vec_supported(),
    reason="sqlite3 loadable extensions unavailable on this Python build",
)
# ---------------------------------------------------------------------------
# parse_document
# ---------------------------------------------------------------------------


def _tiny_pdf_bytes(text: str) -> bytes:
    """Synthesize a real one-page PDF containing ``text`` via pypdf."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 24 Tf 100 700 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    font = DictionaryObject()
    font[NameObject("/Type")] = NameObject("/Font")
    font[NameObject("/Subtype")] = NameObject("/Type1")
    font[NameObject("/BaseFont")] = NameObject("/Helvetica")
    resources = DictionaryObject()
    resources[NameObject("/Font")] = DictionaryObject(
        {NameObject("/F1"): writer._add_object(font)}
    )
    page[NameObject("/Resources")] = resources
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _tiny_docx_bytes(text: str) -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_parse_document_markdown_and_txt():
    assert parse_document("notes.md", b"# Title\n\nBody text.") == "# Title\n\nBody text."
    assert parse_document("notes.txt", "café".encode()) == "café"


def test_parse_document_pdf():
    data = _tiny_pdf_bytes("Hello Phase Four")
    assert "Hello Phase Four" in parse_document("paper.pdf", data)


def test_parse_document_docx():
    data = _tiny_docx_bytes("docx paragraph content")
    assert "docx paragraph content" in parse_document("report.docx", data)


def test_parse_document_rejects_unsupported():
    with pytest.raises(ValueError, match="Unsupported"):
        parse_document("evil.exe", b"MZ")


# ---------------------------------------------------------------------------
# chunk_text
# ---------------------------------------------------------------------------


def test_chunk_text_empty():
    assert chunk_text("") == []
    assert chunk_text("   \n\t  ") == []


def test_chunk_text_short_single_chunk():
    chunks = chunk_text("alpha beta gamma", chunk_size=800, overlap=120)
    assert chunks == [{"text": "alpha beta gamma", "start": 0, "end": 16}]


def test_chunk_text_offsets_exact_and_overlap():
    text = "First sentence here. Second sentence follows. Third and final one."
    chunks = chunk_text(text, chunk_size=30, overlap=10)
    assert len(chunks) >= 2
    joined_start = 0
    for chunk in chunks:
        assert text[chunk["start"] : chunk["end"]] == chunk["text"]
        assert chunk["end"] - chunk["start"] <= 30
        assert chunk["start"] >= joined_start
    # Overlap: the next chunk starts at or before the previous end.
    for prev, cur in zip(chunks, chunks[1:]):
        assert cur["start"] <= prev["end"]
    # Coverage: last chunk reaches the end of the input.
    assert chunks[-1]["end"] == len(text.rstrip())


def test_chunk_text_hard_splits_overlong_word():
    word = "x" * 100
    chunks = chunk_text(word, chunk_size=30, overlap=0)
    assert [c["text"] for c in chunks] == ["x" * 30, "x" * 30, "x" * 30, "x" * 10]
    assert [c["start"] for c in chunks] == [0, 30, 60, 90]
    assert [c["end"] for c in chunks] == [30, 60, 90, 100]


def test_chunk_text_sentence_boundary():
    text = "alpha beta gamma delta. epsilon " + "padding " * 40
    chunks = chunk_text(text, chunk_size=50, overlap=0)
    assert chunks[0]["text"] == "alpha beta gamma delta."
    assert chunks[0]["end"] == 23


# ---------------------------------------------------------------------------
# Upload endpoint
# ---------------------------------------------------------------------------


def _create_workflow() -> str:
    response = client.post(
        "/api/workflows",
        json={
            "name": "RAG upload",
            "graph_json": valid_graph(
                [
                    {
                        "id": "n1",
                        "position": {"x": 0, "y": 0},
                        "data": {"label": "Calc", "nodeType": "tool", "toolType": "calculator"},
                    }
                ],
            ),
        },
    )
    assert response.status_code == 200
    return response.json()["id"]


def _md_body(words: int) -> bytes:
    # ~6 chars/word; words=400 → ~2400 chars > 800 → multi-chunk
    return ("alpha beta gamma delta epsilon zeta " * ((words // 6) + 1))[: words * 6].encode()


def test_upload_markdown_creates_chunks(monkeypatch):
    monkeypatch.setattr("app.services.embeddings.embed_text", _hashing_vector)
    workflow_id = _create_workflow()
    response = client.post(
        f"/api/workflows/{workflow_id}/knowledge/upload",
        files={"file": ("notes.md", _md_body(400), "text/markdown")},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["chunks_created"] > 1
    assert payload["chunk_size"] == 800
    assert payload["overlap"] == 120
    assert payload["filename"] == "notes.md"
    assert UUID(payload["document_id"])

    docs = client.get(f"/api/workflows/{workflow_id}/knowledge").json()
    assert len(docs) == payload["chunks_created"]
    for i, doc in enumerate(sorted(docs, key=lambda d: d["chunk_index"])):
        assert doc["source_document_id"] == payload["document_id"]
        assert doc["source_name"] == "notes.md"
        assert doc["chunk_index"] == i
        assert doc["title"] == f"notes.md#{i + 1}"
        assert doc["chunk_start"] < doc["chunk_end"]


def test_upload_rejects_unsupported_type():
    workflow_id = _create_workflow()
    response = client.post(
        f"/api/workflows/{workflow_id}/knowledge/upload",
        files={"file": ("payload.exe", b"MZ" + b"\x00" * 100, "application/x-msdownload")},
    )
    assert response.status_code == 400


def test_upload_rejects_oversized():
    workflow_id = _create_workflow()
    big = b"a" * (8 * 1024 * 1024 + 1)
    response = client.post(
        f"/api/workflows/{workflow_id}/knowledge/upload",
        files={"file": ("big.txt", big, "text/plain")},
    )
    assert response.status_code == 413


# ---------------------------------------------------------------------------
# sqlite-vec store
# ---------------------------------------------------------------------------


def _make_workflow_row(db) -> models.Workflow:
    row = models.Workflow(user_id=uuid.uuid4(), name=f"vec-{uuid.uuid4()}")
    db.add(row)
    db.flush()
    return row


@vec_required
def test_sqlite_vec_store_search_roundtrip():
    from app.services.vector_stores import SqliteVecStore, get_vector_store

    db = SessionLocal()
    try:
        wf = _make_workflow_row(db)
        other = _make_workflow_row(db)
        near = models.KnowledgeDocument(workflow_id=wf.id, title="near", text="a")
        far = models.KnowledgeDocument(workflow_id=wf.id, title="far", text="b")
        noise = models.KnowledgeDocument(workflow_id=other.id, title="noise", text="c")
        db.add_all([near, far, noise])
        db.flush()

        store = get_vector_store(db)
        assert isinstance(store, SqliteVecStore)
        query = _hashing_vector("alpha beta gamma")
        store.store(db, near.id, _hashing_vector("alpha beta gamma epsilon"))
        store.store(db, far.id, _hashing_vector("zebra yak xylophone"))
        # Nearest overall, but must be excluded by the workflow filter.
        store.store(db, noise.id, query)

        hits = store.search(db, query, wf.id, top_k=2)
        assert [h["document_id"] for h in hits] == [str(near.id), str(far.id)]
        assert hits[0]["distance"] < hits[1]["distance"]

        store.delete(db, far.id)
        hits = store.search(db, query, wf.id, top_k=2)
        assert [h["document_id"] for h in hits] == [str(near.id)]
        db.rollback()
    finally:
        db.close()

@vec_required
def test_retrieve_merges_legacy_json_corpus(monkeypatch):
    """A doc created before the vector store existed (JSON embedding only, not
    indexed) must still be retrievable: retrieve_documents merges store hits
    with JSON-scored unindexed docs so the first indexed hit no longer hides
    the legacy corpus — read-only, no writes to the shared index."""
    monkeypatch.setattr("app.config.settings.vector_store", "sqlite_vec")
    from app.services.vector_search import retrieve_documents
    from app.services.vector_stores import get_vector_store

    db = SessionLocal()
    try:
        wf = _make_workflow_row(db)
        legacy_id = uuid.uuid4()
        new_id = uuid.uuid4()
        legacy_vec = _hashing_vector("alpha beta gamma legacy")
        new_vec = _hashing_vector("zebra yak xylophone new")
        db.add(models.KnowledgeDocument(id=legacy_id, workflow_id=wf.id, title="legacy", text="legacy doc"))
        db.add(models.KnowledgeDocument(id=new_id, workflow_id=wf.id, title="new", text="new doc"))
        db.commit()
        # Only the NEW doc is in the vec index — the legacy row predates it.
        store = get_vector_store(db)
        store.store(db, new_id, new_vec)
        db.commit()

        monkeypatch.setattr("app.services.embeddings.embed_text", lambda _q: legacy_vec)
        docs = [
            {"id": str(legacy_id), "title": "legacy", "text": "legacy doc", "embedding": legacy_vec},
            {"id": str(new_id), "title": "new", "text": "new doc", "embedding": new_vec},
        ]
        hits = retrieve_documents(wf.id, "alpha beta gamma legacy", docs, top_k=2)
        ids = {h["id"] for h in hits}
        assert str(legacy_id) in ids, "legacy JSON-only doc must be retrievable via merge"
    finally:
        db.close()


# ---------------------------------------------------------------------------
# kb_retrieve citations
# ---------------------------------------------------------------------------


def test_kb_retrieve_includes_citations_for_workflow_source(monkeypatch):
    # JSON-embedding fallback path is what runs under `vector_store=none`;
    # bm25 needs no embeddings at all.
    monkeypatch.setattr("app.config.settings.vector_store", "none")
    db = SessionLocal()
    try:
        wf = _make_workflow_row(db)
        source_id = uuid.uuid4()
        db.add(
            models.KnowledgeDocument(
                workflow_id=wf.id,
                title="runbook.md#1",
                text="restart the aegis worker process to recover",
                source_document_id=source_id,
                source_name="runbook.md",
                source_mime="text/markdown",
                chunk_index=0,
                chunk_start=5,
                chunk_end=52,
            )
        )
        db.add(
            models.KnowledgeDocument(
                workflow_id=wf.id,
                title="other",
                text="unrelated dessert recipes",
            )
        )
        db.commit()

        ctx = {
            "input": {"q": "aegis worker restart"},
            "steps": {},
            "last_output": "",
            "memory": {},
            "_workflow_id": str(wf.id),
        }
        fn = _make_kb_retrieve_fn(
            "kb", "{{input.q}}", [], 3, "kb_retrieve", ctx, kb_source="workflow"
        )
        out = json.loads(fn(""))
    finally:
        db.close()

    assert out["source"] == "workflow"
    citations = out["citations"]
    assert citations and citations[0]["document_id"]
    top = citations[0]
    assert top["rank"] == 1
    assert top["title"] == "runbook.md#1"
    assert top["source_name"] == "runbook.md"
    assert top["source_document_id"] == str(source_id)
    assert top["chunk_index"] == 0
    assert top["char_start"] == 5
    assert top["char_end"] == 52
    assert top["score"] > 0


def test_kb_retrieve_citations_null_span_for_inline_docs():
    ctx = {"input": {"q": "python programming"}, "steps": {}, "last_output": "", "memory": {}}
    docs = [{"id": "a", "text": "python programming language guide", "title": "py"}]
    fn = _make_kb_retrieve_fn("kb", "{{input.q}}", docs, 2, "kb_retrieve", ctx)
    out = json.loads(fn(""))
    assert "citations" in out
    assert out["citations"][0]["source_name"] is None
    assert out["citations"][0]["char_start"] is None
    assert out["citations"][0]["chunk_index"] is None
